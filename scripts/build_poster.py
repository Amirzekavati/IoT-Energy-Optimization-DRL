"""
Fill the Lockwood 4:3 poster template with the latest run results and charts.

Usage:
    python scripts/build_poster.py
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    r"c:\Users\amirz\Downloads\PosterPresentations.com-Virtual(4.3-Ratio)-Template-Lockwood (1).pptx"
)
OUTPUT = ROOT / "docs" / "poster" / "IoT_Energy_Optimization_Poster.pptx"
OUTPUT_FALLBACK = ROOT / "docs" / "poster" / "IoT_Energy_Optimization_Poster_updated.pptx"

CHARTS = ROOT / "assets" / "charts"
ARCH = ROOT / "assets" / "architecture" / "high_level_architecture.png"
METRICS_JSON = CHARTS / "poster_metrics.json"

CHART_CMP = CHARTS / "policy_comparison.png"
CHART_ENERGY = CHARTS / "energy_curves.png"
CHART_AOI = CHARTS / "aoi_curves.png"
CHART_PACKETS = CHARTS / "packets_curves.png"

INK = RGBColor(0x1A, 0x23, 0x32)
ACCENT = RGBColor(0x0F, 0x4C, 0x81)


def _set_runs(paragraph, text, size_pt, bold=False, color=INK, font_name="Tahoma"):
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name


def set_block(shape, title_or_lines, *, size_pt=16, bold=False, align=PP_ALIGN.RIGHT, color=INK):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    lines = title_or_lines if isinstance(title_or_lines, list) else [title_or_lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = 0
        _set_runs(p, line, size_pt=size_pt, bold=bold, color=color)


def set_body_paragraphs(shape, paragraphs, *, size_pt=13, align=PP_ALIGN.RIGHT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, block in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.level = 0
        p.space_after = Pt(6)
        _set_runs(p, block, size_pt=size_pt, bold=False, color=INK)


def _fmt(v, digits=1):
    if v is None:
        return "-"
    try:
        return f"{float(v):.{digits}f}"
    except (TypeError, ValueError):
        return str(v)


def load_metrics():
    if METRICS_JSON.exists():
        return json.loads(METRICS_JSON.read_text(encoding="utf-8"))
    return {}


def fit_picture(slide, path, left, top, box_w, box_h):
    pic = slide.shapes.add_picture(str(path), left, top, width=box_w)
    if pic.height > box_h:
        ratio = box_h / pic.height
        pic.height = box_h
        pic.width = int(pic.width * ratio)
        pic.left = int(left + (box_w - pic.width) / 2)
    if pic.width > box_w:
        ratio = box_w / pic.width
        pic.width = box_w
        pic.height = int(pic.height * ratio)
    if pic.height < box_h:
        pic.top = int(top + (box_h - pic.height) / 2)
    if pic.width < box_w:
        pic.left = int(left + (box_w - pic.width) / 2)
    return pic


def build():
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")

    metrics = load_metrics()
    tx = metrics.get("always_transmit", {})
    sleep = metrics.get("always_sleep", {})
    rnd = metrics.get("random", {})
    dqn = metrics.get("dqn", {})

    prs = Presentation(str(TEMPLATE))
    slide = prs.slides[0]
    shapes = list(slide.shapes)

    body_tl, title_tl = shapes[0], shapes[1]
    body_bl, title_bl = shapes[2], shapes[3]
    body_tm, title_tm = shapes[4], shapes[5]
    title_tr, body_tr = shapes[6], shapes[7]
    title_br, body_br = shapes[8], shapes[9]
    meta, subtitle, title = shapes[10], shapes[11], shapes[12]

    set_block(
        title,
        "بهینه‌سازی مصرف انرژی در شبکه‌های اینترنت اشیا با یادگیری تقویتی عمیق",
        size_pt=28,
        bold=True,
        align=PP_ALIGN.CENTER,
        color=ACCENT,
    )
    set_block(
        subtitle,
        "IoT Energy Optimization using Deep Reinforcement Learning (DQN)",
        size_pt=16,
        bold=False,
        align=PP_ALIGN.CENTER,
        color=INK,
    )
    set_block(
        meta,
        "دانشجو: امیرمحمد ذکاوتی  |  استاد راهنما: دکتر منیره عبدوس  |  دانشگاه شهید بهشتی  |  ۱۴۰۴",
        size_pt=12,
        bold=False,
        align=PP_ALIGN.CENTER,
        color=INK,
    )

    set_block(title_tl, "مسئله و انگیزه", size_pt=18, bold=True, align=PP_ALIGN.RIGHT, color=ACCENT)
    set_body_paragraphs(
        body_tl,
        [
            "نودهای حسگر IoT با باتری محدود کار می‌کنند. ارسال مداوم عمر شبکه را کم و خواب بیش از حد کیفیت پایش را پایین می‌آورد.",
            "هدف: تصمیم هوشمند Transmit/Sleep برای تعادل بین انرژی، تحویل بسته و تازگی داده (AoI).",
        ],
        size_pt=12,
    )

    set_block(title_tm, "اهداف و معماری", size_pt=18, bold=True, align=PP_ALIGN.RIGHT, color=ACCENT)
    set_body_paragraphs(
        body_tm,
        [
            "• شبیه‌ساز IoT + انرژی فاصله‌آگاه + افت بسته",
            "• محیط Gymnasium و آموزش DQN",
            "• مقایسه با Transmit / Sleep / Random",
            "• داشبورد Streamlit",
            "معماری: UI → Simulation → RL → Analytics",
        ],
        size_pt=12,
    )

    set_block(title_tr, "روش پیشنهادی", size_pt=18, bold=True, align=PP_ALIGN.RIGHT, color=ACCENT)
    set_body_paragraphs(
        body_tr,
        [
            "• هزینه TX وابسته به فاصله تا Gateway",
            "• کانال ساده: احتمال تحویل با فاصله کم می‌شود",
            "• معیار AoI برای تازگی اطلاعات",
            "• پاداش: بسته/عمر بیشتر − انرژی − AoI − drop",
            "• ابزار: Python، Gymnasium، SB3، Streamlit",
        ],
        size_pt=12,
    )

    set_block(title_bl, "نتایج کلیدی (ران نهایی)", size_pt=18, bold=True, align=PP_ALIGN.RIGHT, color=ACCENT)
    set_body_paragraphs(
        body_bl,
        [
            "تنظیمات ران: ۳ نود | افق ۵۰ | آموزش ۶۰۰۰ | میانگین ۵ اپیزود | seed=42",
            (
                f"Transmit → Life {_fmt(tx.get('lifetime_steps'))} | "
                f"Pkt {_fmt(tx.get('packets_received'))} | "
                f"PDR {_fmt(tx.get('packet_delivery_ratio'), 2)}"
            ),
            (
                f"Sleep → Life {_fmt(sleep.get('lifetime_steps'))} | "
                f"Pkt {_fmt(sleep.get('packets_received'))} | "
                f"AoI {_fmt(sleep.get('mean_aoi'))}"
            ),
            (
                f"Random → Life {_fmt(rnd.get('lifetime_steps'))} | "
                f"Pkt {_fmt(rnd.get('packets_received'))} | "
                f"PDR {_fmt(rnd.get('packet_delivery_ratio'), 2)}"
            ),
            (
                f"DQN → Life {_fmt(dqn.get('lifetime_steps'))} | "
                f"Pkt {_fmt(dqn.get('packets_received'))} | "
                f"PDR {_fmt(dqn.get('packet_delivery_ratio'), 2)} | "
                f"AoI {_fmt(dqn.get('mean_aoi'))}"
            ),
            "پیام: DQN بهترین مصالحه عمر شبکه، تحویل بسته و تازگی داده را دارد.",
        ],
        size_pt=11,
    )

    set_block(title_br, "نتیجه‌گیری و مراجع", size_pt=18, bold=True, align=PP_ALIGN.RIGHT, color=ACCENT)
    set_body_paragraphs(
        body_br,
        [
            "با مدل فاصله/افت بسته/AoI، DQN از سیاست‌های ثابت بهتر مصالحه می‌کند.",
            "آینده: آموزش بهتر برای n بزرگ‌تر، Multi-Agent RL، سخت‌افزار واقعی.",
            "[1] Transforma Insights, 2026",
            "[4] Heinzelman et al., LEACH, 2000",
            "[8] Banerjee et al., NashDQNSleep, 2025",
            "[9] Leong et al., DRL Sensor Scheduling, 2018",
        ],
        size_pt=11,
    )

    # Center graphics: architecture on top strip + 2x2 charts
    left_x = Emu(900_000)
    usable_w = prs.slide_width - Emu(1_800_000)
    gap = Emu(180_000)

    arch_top = Emu(7_150_000)
    arch_h = Emu(2_700_000)
    if ARCH.exists():
        fit_picture(slide, ARCH, left_x, arch_top, usable_w, arch_h)

    grid_top = arch_top + arch_h + gap
    grid_h = Emu(6_400_000)
    cell_w = int((usable_w - gap) / 2)
    cell_h = int((grid_h - gap) / 2)

    grid_images = [
        (CHART_CMP, 0, 0),
        (CHART_ENERGY, 1, 0),
        (CHART_AOI, 0, 1),
        (CHART_PACKETS, 1, 1),
    ]
    for path, col, row in grid_images:
        if not path.exists():
            raise FileNotFoundError(path)
        x = left_x + col * (cell_w + gap)
        y = grid_top + row * (cell_h + gap)
        fit_picture(slide, path, x, y, cell_w, cell_h)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(OUTPUT))
        print(f"Saved: {OUTPUT}")
    except PermissionError:
        prs.save(str(OUTPUT_FALLBACK))
        print(f"Main poster file is locked. Saved: {OUTPUT_FALLBACK}")


if __name__ == "__main__":
    build()
